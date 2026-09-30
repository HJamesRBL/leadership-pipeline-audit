#!/usr/bin/env python3
"""
Builds the two manual (paper) versions of the Leadership Pipeline Audit:

  1. Leadership Contribution Self-Assessment.docx
  2. Leadership Contribution Assessment - Manager or Peer.docx

Both are the same instrument; only the point of view changes (first person vs.
third person). Run with a python that has python-docx installed:

    python build_assessments.py
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT_DIR = Path(__file__).parent

# --- RBL brand palette (see "RBL color reference slide") ----------------------
NAVY = "071D49"
BLUE = "0086D6"
RED = "ED1B34"
GRAY_LINE = "CCCCCC"
GRAY_FILL = "F2F4F7"
GRAY_TEXT = "666666"

BODY_FONT = "Open Sans"
DISPLAY_FONT = "Georgia"

CHECKBOX = "☐"  # ballot box

# -----------------------------------------------------------------------------
# The instrument.
#
# Ten dimensions drawn from the Four Stages descriptors in the audit deck and
# the app's contribution categories. Each dimension is one forced-choice item
# with four anchors spanning the same spectrum, ordered A -> D.
#
# Deliberately NOT labelled with stage numbers anywhere on the form: naming
# stages invites the rater to equate stage with tenure, title, or age.
#
# Each anchor is (self-assessment wording, manager/peer wording).
# -----------------------------------------------------------------------------
ITEMS = [
    (
        "Where the direction for my work comes from",
        "Where the direction for their work comes from",
        [
            ("I am told what to work on and check in often to stay on track.",
             "They are told what to work on and check in often to stay on track."),
            ("I am given a goal and decide for myself how to get there.",
             "They are given a goal and decide for themselves how to get there."),
            ("I set direction for a team and tie their work to broader priorities.",
             "They set direction for a team and tie its work to broader priorities."),
            ("I set direction for a business, or a major part of one.",
             "They set direction for a business, or a major part of one."),
        ],
    ),
    (
        "The scope of what I am accountable for",
        "The scope of what they are accountable for",
        [
            ("A defined piece of a project that someone else owns.",
             "A defined piece of a project that someone else owns."),
            ("Whole projects that I own from start to finish.",
             "Whole projects that they own from start to finish."),
            ("The output of a group, integrated across several areas of work.",
             "The output of a group, integrated across several areas of work."),
            ("Which work the organization takes on in the first place.",
             "Which work the organization takes on in the first place."),
        ],
    ),
    (
        "What people value me for",
        "What people value this person for",
        [
            ("Being dependable — I deliver the basics and I am easy to work with.",
             "Being dependable — they deliver the basics and are easy to work with."),
            ("My expertise — people come to me because I know this area best.",
             "Their expertise — people go to them because they know the area best."),
            ("Making others effective and raising the capability of the group.",
             "Making others effective and raising the capability of the group."),
            ("My judgment on where the organization should go and what it should stop.",
             "Their judgment on where the organization should go and what it should stop."),
        ],
    ),
    (
        "How my results actually get produced",
        "How their results actually get produced",
        [
            ("Through my own effort, with regular guidance from someone more experienced.",
             "Through their own effort, with regular guidance from someone more experienced."),
            ("Through my own effort, working independently.",
             "Through their own effort, working independently."),
            ("Through other people — by enabling, coaching, and influencing them.",
             "Through other people — by enabling, coaching, and influencing them."),
            ("Through structure and resources — I shape the conditions others work in.",
             "Through structure and resources — they shape the conditions others work in."),
        ],
    ),
    (
        "My part in other people's development",
        "Their part in other people's development",
        [
            ("I am mostly the one learning, from more experienced colleagues.",
             "They are mostly the one learning, from more experienced colleagues."),
            ("I share what I know when asked, but developing others is not really my job.",
             "They share what they know when asked, but developing others is not their focus."),
            ("I coach, mentor, and develop people as a core part of my work.",
             "They coach, mentor, and develop people as a core part of their work."),
            ("I sponsor people into key roles and test them for bigger responsibility.",
             "They sponsor people into key roles and test them for bigger responsibility."),
        ],
    ),
    (
        "Who I work with regularly",
        "Who they work with regularly",
        [
            ("My immediate team and the people who direct my work.",
             "Their immediate team and the people who direct their work."),
            ("A network of specialists in my own discipline.",
             "A network of specialists in their own discipline."),
            ("A broad internal network across functions, plus contacts in my industry.",
             "A broad internal network across functions, plus contacts in the industry."),
            ("External stakeholders — customers, investors, partners, the market.",
             "External stakeholders — customers, investors, partners, the market."),
        ],
    ),
    (
        "How I get others to act",
        "How they get others to act",
        [
            ("I mostly accept direction; I have little say beyond my own work.",
             "They mostly accept direction and have little say beyond their own work."),
            ("I persuade people on the merits of my technical judgment.",
             "They persuade people on the merits of their technical judgment."),
            ("I use formal and informal influence to move ideas through the organization.",
             "They use formal and informal influence to move ideas through the organization."),
            ("I control resources — people, money, information — and decide where they go.",
             "They control resources — people, money, information — and decide where they go."),
        ],
    ),
    (
        "What I am typically thinking about",
        "What they are typically thinking about",
        [
            ("What is in front of me this week.",
             "What is in front of them this week."),
            ("The problems and standards of my own discipline.",
             "The problems and standards of their own discipline."),
            ("How my group's work fits the bigger picture and connects to other groups.",
             "How their group's work fits the bigger picture and connects to other groups."),
            ("Where the whole organization needs to be several years from now.",
             "Where the whole organization needs to be several years from now."),
        ],
    ),
    (
        "What I do when work hits an obstacle",
        "What they do when work hits an obstacle",
        [
            ("I raise it and wait for direction on how to handle it.",
             "They raise it and wait for direction on how to handle it."),
            ("I find a way around it myself and deliver anyway.",
             "They find a way around it themselves and deliver anyway."),
            ("I clear barriers out of the way so other people's work can move.",
             "They clear barriers out of the way so other people's work can move."),
            ("I change the systems and processes that create the barrier.",
             "They change the systems and processes that create the barrier."),
        ],
    ),
    (
        "Where my new ideas go",
        "Where their new ideas go",
        [
            ("I take initiative inside the boundaries I am given.",
             "They take initiative inside the boundaries they are given."),
            ("I innovate within my own area of expertise.",
             "They innovate within their own area of expertise."),
            ("I carry ideas across boundaries and get other groups behind them.",
             "They carry ideas across boundaries and get other groups behind them."),
            ("I set the agenda for what the organization takes on next.",
             "They set the agenda for what the organization takes on next."),
        ],
    ),
]

LETTERS = ["A", "B", "C", "D"]


# -----------------------------------------------------------------------------
# Low-level docx helpers
# -----------------------------------------------------------------------------
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def cell_borders(cell, **edges):
    """edges e.g. top={'sz': 6, 'color': 'CCCCCC'} ; omitted edges -> none."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        spec = edges.get(edge)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), spec.get("val", "single"))
            el.set(qn("w:sz"), str(spec.get("sz", 6)))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), spec.get("color", GRAY_LINE))
        borders.append(el)
    tc_pr.append(borders)


def cell_margins(cell, top=0, bottom=0, left=60, right=60):
    """Margins in twentieths of a point."""
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def row_height(row, points, exact=True):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(points * 20)))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    tr_pr.append(h)


def keep_together(paragraph, with_next=True):
    """keepLines stops a paragraph splitting mid-block. keepNext chains it to the
    paragraph that follows, so use it sparingly — a long chain of keepNext
    paragraphs is unbreakable and will shove whole sections onto a new page."""
    p_pr = paragraph._p.get_or_add_pPr()
    tags = ("w:keepNext", "w:keepLines") if with_next else ("w:keepLines",)
    for tag in tags:
        p_pr.append(OxmlElement(tag))


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def no_table_borders(table):
    """Clear the table-level grid; per-cell borders still apply."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    table._tbl.tblPr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)


def run(paragraph, text, *, size=9, bold=False, italic=False,
        color=None, font=BODY_FONT):
    r = paragraph.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    rpr = r._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), font)
    return r


def para(container, text="", *, size=9, bold=False, italic=False, color=None,
         font=BODY_FONT, space_before=0, space_after=0, align=None,
         line_spacing=1.08):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    if text:
        run(p, text, size=size, bold=bold, italic=italic, color=color, font=font)
    return p


def cell_para(cell, text="", **kw):
    """First paragraph of a cell is reused; later ones are appended."""
    if cell.paragraphs and not cell.paragraphs[0].runs and not getattr(cell, "_used", False):
        p = cell.paragraphs[0]
        cell._used = True
        p.paragraph_format.space_before = Pt(kw.get("space_before", 0))
        p.paragraph_format.space_after = Pt(kw.get("space_after", 0))
        p.paragraph_format.line_spacing = kw.get("line_spacing", 1.08)
        if kw.get("align") is not None:
            p.alignment = kw["align"]
        if text:
            run(p, text, size=kw.get("size", 9), bold=kw.get("bold", False),
                italic=kw.get("italic", False), color=kw.get("color"),
                font=kw.get("font", BODY_FONT))
        return p
    return para(cell, text, **kw)


def spacer(doc, points):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run(p, "", size=points)
    return p


NBSP = " "  # ordinary trailing spaces lose their underline when rendered


def underline_field(paragraph, label, blank_chars=18, size=9):
    if label:
        run(paragraph, label, size=size, bold=True, color=NAVY)
    r = run(paragraph, NBSP * blank_chars, size=size)
    r.underline = True


# w:tblPr children must appear in schema order, so new elements get inserted
# rather than appended.
_TBLPR_AFTER_W = ("w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                  "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                  "w:tblDescription", "w:tblPrChange")
_TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                        "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def set_table_width(table, inches):
    """python-docx sizes tables from cell widths alone, which drifts once cell
    margins are involved; pin the table itself so stacked bands line up."""
    table.autofit = False  # emits w:tblLayout fixed
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(int(inches * 1440)))
    w.set(qn("w:type"), "dxa")
    table._tbl.tblPr.insert_element_before(w, *_TBLPR_AFTER_W)


# -----------------------------------------------------------------------------
# Document sections
# -----------------------------------------------------------------------------
def setup_page(doc):
    s = doc.sections[0]
    s.page_width = Inches(8.5)
    s.page_height = Inches(11)
    s.left_margin = Inches(0.55)
    s.right_margin = Inches(0.55)
    s.top_margin = Inches(0.45)
    s.bottom_margin = Inches(0.4)
    s.footer_distance = Inches(0.25)

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(9)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), BODY_FONT)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.08

    footer_p = s.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(footer_p, "© The RBL Group   |   Leadership Pipeline Audit",
        size=7.5, color=GRAY_TEXT)


def add_masthead(doc, title, subtitle):
    """Navy title band with the RBL red rule beneath it. Both live in one table
    so the rule can never end up wider or narrower than the band above it."""
    t = doc.add_table(rows=2, cols=1)
    set_table_width(t, 7.4)

    band = t.cell(0, 0)
    band.width = Inches(7.4)
    shade(band, NAVY)
    cell_borders(band)
    cell_margins(band, top=110, bottom=110, left=170, right=170)
    band.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = cell_para(band, space_after=1)
    run(p, title, size=17, font=DISPLAY_FONT, color="FFFFFF")
    p2 = cell_para(band, space_before=0)
    run(p2, subtitle, size=8.5, color="AFC6E0")

    accent = t.cell(1, 0)
    accent.width = Inches(7.4)
    shade(accent, RED)
    cell_borders(accent)
    cell_margins(accent, 0, 0, 0, 0)
    row_height(t.rows[1], 3.5)
    ap = cell_para(accent, line_spacing=1.0)
    run(ap, "", size=1)

    no_table_borders(t)


def add_field_line(doc, fields):
    spacer(doc, 5)
    p = para(doc, space_after=0)
    for i, (label, blanks) in enumerate(fields):
        if i:
            run(p, "     ", size=9)
        underline_field(p, label, blanks)


def add_instruction_box(doc, body):
    spacer(doc, 5)
    t = doc.add_table(rows=1, cols=1)
    set_table_width(t, 7.4)
    c = t.cell(0, 0)
    c.width = Inches(7.4)
    shade(c, GRAY_FILL)
    cell_borders(c, left={"sz": 18, "color": BLUE})
    cell_margins(c, top=90, bottom=90, left=150, right=150)
    p = cell_para(c, line_spacing=1.15)
    run(p, body, size=8.8, color="333333")
    no_table_borders(t)


def add_section_head(doc, kicker, heading, space_before=9):
    spacer(doc, space_before)
    p = para(doc, space_after=1)
    keep_together(p)
    run(p, kicker.upper(), size=8, bold=True, color=BLUE)
    p2 = para(doc, space_after=3)
    keep_together(p2)
    run(p2, heading, size=12.5, font=DISPLAY_FONT, color=NAVY)


def add_item_table(doc, third_person):
    t = doc.add_table(rows=0, cols=4)
    set_table_width(t, 7.4)
    col_w = Inches(7.4 / 4)

    for idx, (self_q, other_q, anchors) in enumerate(ITEMS, start=1):
        question = other_q if third_person else self_q

        q_row = t.add_row()
        cant_split(q_row)
        for c in q_row.cells:
            c.width = col_w
        qcell = q_row.cells[0].merge(q_row.cells[3])
        shade(qcell, "FFFFFF")
        cell_borders(qcell, top={"sz": 6, "color": GRAY_LINE})
        cell_margins(qcell, top=60, bottom=15, left=0, right=0)
        qp = cell_para(qcell, space_after=0)
        keep_together(qp)  # question stays glued to its four options
        run(qp, f"{idx}.  ", size=9.5, bold=True, color=BLUE)
        run(qp, question, size=9.5, bold=True, color=NAVY)

        o_row = t.add_row()
        cant_split(o_row)
        for j, cell in enumerate(o_row.cells):
            cell.width = col_w
            cell_borders(cell)
            cell_margins(cell, top=10, bottom=45, left=0, right=100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            text = anchors[j][1] if third_person else anchors[j][0]
            p = cell_para(cell, space_after=0, line_spacing=1.02)
            keep_together(p, with_next=False)
            run(p, f"{CHECKBOX} ", size=10, color=NAVY)
            run(p, f"{LETTERS[j]}  ", size=8.4, bold=True, color=GRAY_TEXT)
            run(p, text, size=8.4, color="333333")

    no_table_borders(t)
    return t


def add_performance_scale(doc, third_person):
    lead = (
        "Set aside how this person contributes for a moment. Compared with their peers "
        "— people at a similar level doing comparable work — how would you rate "
        "their overall performance and impact? Check one number."
        if third_person else
        "Set aside how you contribute for a moment. Compared with your peers — people "
        "at a similar level doing comparable work — how would you rate your overall "
        "performance and impact? Check one number."
    )
    para(doc, lead, size=8.8, space_after=5, line_spacing=1.15)

    t = doc.add_table(rows=2, cols=10)
    set_table_width(t, 7.4)
    col_w = Inches(7.4 / 10)
    box = {"sz": 6, "color": "999999"}

    for i in range(10):
        c = t.cell(0, i)
        c.width = col_w
        cell_borders(c, top=box, left=box, bottom=box, right=box)
        cell_margins(c, top=40, bottom=40, left=0, right=0)
        c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell_para(c, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        run(p, str(i + 1), size=11, bold=True, color=NAVY)

        b = t.cell(1, i)
        b.width = col_w
        cell_borders(b)
        cell_margins(b, top=30, bottom=0, left=0, right=0)
        p = cell_para(b, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        run(p, CHECKBOX, size=10, color=NAVY)

    no_table_borders(t)

    labels = doc.add_table(rows=1, cols=3)
    set_table_width(labels, 7.4)
    caps = [
        ("1  –  Lowest performance among peers", WD_ALIGN_PARAGRAPH.LEFT, Inches(2.6)),
        ("5  –  About average", WD_ALIGN_PARAGRAPH.CENTER, Inches(2.2)),
        ("10  –  Highest performance among peers", WD_ALIGN_PARAGRAPH.RIGHT, Inches(2.6)),
    ]
    for i, (text, align, width) in enumerate(caps):
        c = labels.cell(0, i)
        c.width = width
        cell_borders(c)
        cell_margins(c, top=20, bottom=0, left=0, right=0)
        p = cell_para(c, align=align, space_after=0)
        run(p, text, size=7.8, italic=True, color=GRAY_TEXT)
    no_table_borders(labels)


def add_write_in(doc, prompt, lines=4):
    p = para(doc, space_before=9, space_after=4)
    keep_together(p)
    run(p, prompt, size=9, bold=True, color=NAVY)
    for _ in range(lines):
        lp = doc.add_paragraph()
        lp.paragraph_format.space_before = Pt(11)
        lp.paragraph_format.space_after = Pt(0)
        lp.paragraph_format.line_spacing = 1.0
        run(lp, "", size=9)
        p_pr = lp._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), GRAY_LINE)
        borders.append(bottom)
        p_pr.append(borders)


def add_tally_box(doc):
    spacer(doc, 14)
    t = doc.add_table(rows=1, cols=1)
    set_table_width(t, 7.4)
    c = t.cell(0, 0)
    c.width = Inches(7.4)
    shade(c, GRAY_FILL)
    cell_borders(c,
                 top={"sz": 6, "color": NAVY}, bottom={"sz": 6, "color": NAVY},
                 left={"sz": 6, "color": NAVY}, right={"sz": 6, "color": NAVY})
    cell_margins(c, top=110, bottom=110, left=160, right=160)

    p = cell_para(c, space_after=4)
    run(p, "PART ONE TALLY", size=8, bold=True, color=BLUE)

    p2 = para(c, space_after=4)
    run(p2, "Go back through Part One and count how many times you chose each letter.",
        size=8.8, color="333333")

    p3 = para(c, space_before=3, space_after=5)
    for letter in LETTERS:
        run(p3, f"{NBSP * 8}{letter}:", size=11, bold=True, color=NAVY)
        underline_field(p3, "", blank_chars=8, size=11)

    p4 = para(c)
    run(p4,
        "There is no score to add up and no letter that is better than another. "
        "The pattern is the point — your facilitator will walk you through what it means.",
        size=8, italic=True, color=GRAY_TEXT)
    no_table_borders(t)


# -----------------------------------------------------------------------------
# Build
# -----------------------------------------------------------------------------
def build(third_person: bool, filename: str):
    doc = Document()
    setup_page(doc)

    if third_person:
        add_masthead(
            doc,
            "Leadership Contribution Assessment",
            "MANAGER / PEER RATING   •   THE RBL GROUP",
        )
        add_field_line(doc, [
            ("Person you are rating:", 24),
            ("Your name:", 20),
            ("Date:", 12),
        ])
        p = para(doc, space_before=6)
        run(p, "Your relationship to this person:  ", size=9, bold=True, color=NAVY)
        run(p, f"{CHECKBOX} Manager     ", size=9.5)
        run(p, f"{CHECKBOX} Peer     ", size=9.5)
        run(p, f"{CHECKBOX} Other: ", size=9.5)
        r = run(p, " " * 16, size=9)
        r.underline = True

        add_instruction_box(doc, (
            "This is not a performance review. It asks how this person contributes at work "
            "— not how well they do their job, and not their level or title. For each of "
            "the ten items below, check the one option that best describes how they actually "
            "work most of the time. Base your answers on what you have directly observed, "
            "not on their title, tenure, or potential. There are no right answers and no "
            "option is better than another. Takes about 10 minutes."
        ))
    else:
        add_masthead(
            doc,
            "Leadership Contribution Self-Assessment",
            "THE RBL GROUP   •   LEADERSHIP PIPELINE AUDIT",
        )
        add_field_line(doc, [
            ("Your name:", 26),
            ("Role:", 26),
            ("Date:", 14),
        ])
        add_instruction_box(doc, (
            "This is not a performance review. It asks how you contribute at work — not "
            "how well you do your job, and not your level or title. For each of the ten items "
            "below, check the one option that best describes how you actually work most of "
            "the time. Answer for what is true today, not for what your title implies, what "
            "you did in a previous role, or what you are working toward. There are no right "
            "answers and no option is better than another. Takes about 10 minutes."
        ))

    add_section_head(doc, "Part One", "How the work gets done", space_before=8)
    add_item_table(doc, third_person)

    add_section_head(doc, "Part Two", "Relative performance", space_before=15)
    add_performance_scale(doc, third_person)

    add_section_head(doc, "Part Three", "Reflection", space_before=15)
    if third_person:
        add_write_in(doc, "Where does this person make their biggest contribution today, "
                          "and who benefits from it?")
        add_write_in(doc, "What single change in how they work would most increase their impact?")
    else:
        add_write_in(doc, "Where do you make your biggest contribution today, "
                          "and who benefits from it?")
        add_write_in(doc, "What single change in how you work would most increase your impact?")

    add_tally_box(doc)

    path = OUT_DIR / filename
    doc.save(path)
    return path


if __name__ == "__main__":
    for tp, name in (
        (False, "Leadership Contribution Self-Assessment.docx"),
        (True, "Leadership Contribution Assessment - Manager or Peer.docx"),
    ):
        print("wrote", build(tp, name))
